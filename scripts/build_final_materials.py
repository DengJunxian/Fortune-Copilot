#!/usr/bin/env python3
"""Build the final Fortune Copilot technical whitepaper DOCX from its Markdown source."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "technical_whitepaper.md"
OUTPUT = ROOT / "output" / "doc" / "wealthtwin_technical_whitepaper.docx"

NAVY = "17324D"
NAVY_LIGHT = "E9F0F5"
RED = "B3262D"
GOLD = "B48A3A"
TEXT = "23313D"
MUTED = "5C6B76"
GRID = "CED8E0"
PALE_RED = "F8ECEC"
WHITE = "FFFFFF"
# Use a preinstalled macOS system font whose Simplified Chinese glyphs remain
# visible in both Word and LibreOffice's headless renderer. Arial Unicode MS is
# present on this machine but LibreOffice 25 can emit blank CJK glyphs for it.
BODY_FONT = "Heiti SC"
CODE_FONT = "Menlo"


def set_east_asia_font(run, name: str) -> None:  # type: ignore[no-untyped-def]
    run.font.name = name
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for script in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{script}"), name)


def set_code_font(run) -> None:  # type: ignore[no-untyped-def]
    """Use a monospaced Latin font while retaining Chinese glyph coverage."""
    run.font.name = CODE_FONT
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), CODE_FONT)
    fonts.set(qn("w:hAnsi"), CODE_FONT)
    fonts.set(qn("w:eastAsia"), BODY_FONT)
    fonts.set(qn("w:cs"), BODY_FONT)


def set_cell_shading(cell, fill: str) -> None:  # type: ignore[no-untyped-def]
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRID, size: str = "6") -> None:  # type: ignore[no-untyped-def]
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:  # type: ignore[no-untyped-def]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("— ")
    set_east_asia_font(run, BODY_FONT)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])
    tail = paragraph.add_run(" —")
    set_east_asia_font(tail, BODY_FONT)


def add_hyperlink(paragraph, text: str, url: str) -> None:  # type: ignore[no-untyped-def]
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run_element = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for script in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{script}"), BODY_FONT)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1E5F8A")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([fonts, color, underline])
    text_element = OxmlElement("w:t")
    text_element.text = text
    run_element.extend([properties, text_element])
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)


INLINE_TOKEN = re.compile(r"(\[[^\]]+\]\([^)]+\)|`[^`]+`|\*\*[^*]+\*\*)")


def add_inline(paragraph, text: str, *, color: str = TEXT, size: float = 10.5) -> None:  # type: ignore[no-untyped-def]
    cursor = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_east_asia_font(run, BODY_FONT)
            run.font.color.rgb = RGBColor.from_string(color)
            run.font.size = Pt(size)
        token = match.group(0)
        if token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()  # type: ignore[union-attr]
            add_hyperlink(paragraph, label, url)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_code_font(run)
            run.font.size = Pt(size - 0.5)
            run.font.color.rgb = RGBColor.from_string(RED)
        else:
            run = paragraph.add_run(token[2:-2])
            set_east_asia_font(run, BODY_FONT)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(color)
            run.font.size = Pt(size)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_east_asia_font(run, BODY_FONT)
        run.font.color.rgb = RGBColor.from_string(color)
        run.font.size = Pt(size)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.15)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.8)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.3

    for name, size, color in (
        ("Title", 28, NAVY),
        ("Heading 1", 20, NAVY),
        ("Heading 2", 15, NAVY),
        ("Heading 3", 12, RED),
    ):
        style = document.styles[name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("智运财富 Fortune Copilot  ·  技术白皮书 1.0")
    set_east_asia_font(run, BODY_FONT)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)


def add_cover(document: Document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(54)

    eyebrow = document.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = eyebrow.add_run("WEALTHTWIN · COMPETITION DELIVERY")
    set_east_asia_font(run, BODY_FONT)
    run.font.size = Pt(10)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(RED)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(20)
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("智运财富 Fortune Copilot")
    set_east_asia_font(run, BODY_FONT)
    run.font.size = Pt(31)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("中国家庭财富数字孪生与智能投顾操作系统\n技术白皮书")
    set_east_asia_font(run, BODY_FONT)
    run.font.size = Pt(19)
    run.font.color.rgb = RGBColor.from_string(TEXT)

    rule = document.add_table(rows=1, cols=2)
    rule.alignment = WD_TABLE_ALIGNMENT.LEFT
    rule.autofit = False
    rule.columns[0].width = Cm(3.2)
    rule.columns[1].width = Cm(12.0)
    set_cell_shading(rule.cell(0, 0), RED)
    set_cell_shading(rule.cell(0, 1), NAVY)
    for cell in rule.rows[0].cells:
        cell.height = Cm(0.12)
        set_cell_border(cell, color=WHITE, size="0")

    metadata = document.add_paragraph()
    metadata.paragraph_format.space_before = Pt(26)
    metadata.paragraph_format.space_after = Pt(22)
    add_inline(
        metadata,
        "软件版本 0.14.0\n材料版本 1.0\n数据日 2026-08-10\n生成日 2026-08-11",
        color=MUTED,
        size=10,
    )

    boundary = document.add_table(rows=1, cols=1)
    boundary.alignment = WD_TABLE_ALIGNMENT.LEFT
    boundary.autofit = False
    boundary.columns[0].width = Cm(15.2)
    cell = boundary.cell(0, 0)
    set_cell_shading(cell, PALE_RED)
    set_cell_border(cell, color="E4BFC1", size="8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline(
        p,
        "竞赛原型，非中国工商银行官方产品。仅使用合成数据、Mock 产品与 Mock 银行接口；不构成投资、法律、税务或保险建议，不承诺本金或收益。",
        color=RED,
        size=10,
    )

    closing = document.add_paragraph()
    closing.paragraph_format.space_before = Pt(62)
    run = closing.add_run("目标先于产品 · 约束先于收益 · 证据先于解释")
    set_east_asia_font(run, BODY_FONT)
    run.font.size = Pt(11)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(GOLD)

    document.add_page_break()


def add_manual_toc(document: Document) -> None:
    title = document.add_heading("目录", level=1)
    title.paragraph_format.space_after = Pt(16)
    chapters = [
        "1. 中国家庭财富管理问题",
        "2. 华衡自主框架：一核四账、三尺、五硬一软、双画像、六阶段",
        "3. 四重购买力评价",
        "4. 数据字典、家庭数字账本与财富数字孪生",
        "5. 确定性财务指标体系",
        "6. 目标规划与动态四账户算法",
        "7. 组合优化、产品映射与适当性",
        "8. 行为金融双画像、实验与干预",
        "9. RAG、知识图谱、多智能体与 LLM 边界",
        "10. 适当性、隐私、安全与审计",
        "11. 三端架构与工商银行落地边界",
        "12. 实验、局限与未来路线",
        "资料与实现索引",
        "结语",
    ]
    for index, chapter in enumerate(chapters):
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.2 if index < 12 else 0.8)
        p.paragraph_format.space_after = Pt(4)
        add_inline(p, chapter, color=NAVY if index < 12 else MUTED, size=10.5)
    document.add_page_break()


def add_heading(document: Document, text: str, level: int, *, first_h2: bool) -> bool:
    if level == 1:
        return first_h2
    if level == 2:
        p = document.add_heading(text, level=1)
        # Page-break-before keeps chapters on fresh pages without creating a
        # blank page when the preceding chapter ends exactly at a page edge.
        p.paragraph_format.page_break_before = first_h2
        p.paragraph_format.space_after = Pt(10)
        return True
    p = document.add_heading(text, level=min(level - 1, 3))
    p.paragraph_format.space_after = Pt(5)
    return first_h2


def add_body_paragraph(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.space_after = Pt(6)
    add_inline(p, text)


def add_list_item(document: Document, text: str, *, number: str | None = None) -> None:
    # Word's built-in List Number style continues numbering across independent
    # Markdown lists. Preserve the source number explicitly so each list starts
    # where the author intended after round-tripping through LibreOffice.
    p = document.add_paragraph(style="List Bullet" if number is None else None)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    add_inline(p, f"{number}. {text}" if number is not None else text)


def add_quote(document: Document, text: str) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Cm(0.15)
    table.columns[1].width = Cm(14.9)
    set_cell_shading(table.cell(0, 0), GOLD)
    set_cell_shading(table.cell(0, 1), "F6F3EA")
    for cell in table.rows[0].cells:
        set_cell_border(cell, color=WHITE, size="0")
    p = table.cell(0, 1).paragraphs[0]
    add_inline(p, text, color=TEXT, size=10.5)


def add_code(document: Document, lines: list[str]) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Cm(15.1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F5F7")
    set_cell_border(cell, color=GRID, size="6")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("\n".join(lines))
    set_code_font(run)
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(TEXT)


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_shading(
                cell,
                NAVY if row_index == 0 else ("F7F9FA" if row_index % 2 == 0 else WHITE),
            )
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(
                p,
                value,
                color=WHITE if row_index == 0 else TEXT,
                size=8.5 if len(rows[0]) >= 4 else 9,
            )
            for run in p.runs:
                run.bold = row_index == 0
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_markdown(document: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    index = 0
    first_h2_seen = False
    in_code = False
    code_lines: list[str] = []
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_lines
        if paragraph_lines:
            add_body_paragraph(
                document, " ".join(line.strip() for line in paragraph_lines)
            )
            paragraph_lines = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            if in_code:
                add_code(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                index += 1
                continue
            first_h2_seen = add_heading(document, text, level, first_h2=first_h2_seen)
            index += 1
            continue

        if (
            stripped.startswith("|")
            and index + 1 < len(lines)
            and re.match(r"^\|?\s*:?-+", lines[index + 1].strip())
        ):
            flush_paragraph()
            table_rows: list[list[str]] = []
            header = [part.strip() for part in stripped.strip("|").split("|")]
            table_rows.append(header)
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append(
                    [
                        part.strip()
                        for part in lines[index].strip().strip("|").split("|")
                    ]
                )
                index += 1
            add_markdown_table(document, table_rows)
            continue

        bullet = re.match(r"^-\s+(.+)$", stripped)
        numbered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if bullet or numbered:
            flush_paragraph()
            if numbered:
                add_list_item(document, numbered.group(2), number=numbered.group(1))
            else:
                add_list_item(document, bullet.group(1))  # type: ignore[union-attr]
            index += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            add_quote(document, stripped.lstrip("> "))
            index += 1
            continue

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        # Metadata above the abstract is already represented on the cover.
        if not first_h2_seen and (
            stripped.startswith(("版本：", "对应软件版本：", "数据日：", "材料状态："))
            or stripped.endswith("  ")
        ):
            index += 1
            continue

        paragraph_lines.append(stripped)
        index += 1

    flush_paragraph()


def set_core_properties(document: Document) -> None:
    props = document.core_properties
    props.title = "智运财富 Fortune Copilot 技术白皮书"
    props.subject = "中国家庭财富数字孪生与智能投顾操作系统竞赛交付"
    props.author = "智运财富 Fortune Copilot 项目组"
    props.keywords = "家庭财富, 数字孪生, 动态四账户, 适当性, 行为金融, 可信AI"
    props.comments = "竞赛原型，合成数据与 Mock 接口；非工商银行官方产品。"


def main() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    document = Document()
    configure_document(document)
    set_core_properties(document)
    add_cover(document)
    add_manual_toc(document)
    parse_markdown(document, markdown)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    main()
